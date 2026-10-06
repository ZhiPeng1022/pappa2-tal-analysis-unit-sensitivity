"""Convert the v3 manuscript Markdown to a submission-ready DOCX file."""

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(r"C:\Users\Elsa\Documents\Codex\2026-09-25\9-2")
SOURCE = ROOT / "submission" / "manuscript" / "PAPPA2_manuscript_draft_v3_TAL_atlas.md"
OUTPUT = ROOT / "submission" / "manuscript" / "PAPPA2_manuscript_v3.docx"
STOP_MARKER = "## Chinese drafting notes"


def add_runs(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def convert():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    in_code = False

    for line in lines:
        if line.strip().startswith(STOP_MARKER):
            break
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            paragraph = document.add_paragraph(line)
            paragraph.style = document.styles["No Spacing"]
            for run in paragraph.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.strip().startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, line.strip()[2:])
        elif re.match(r"^\d+\.\s", line.strip()):
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+\.\s", "", line.strip()))
        elif line.strip():
            paragraph = document.add_paragraph()
            add_runs(paragraph, line.strip())

    document.save(OUTPUT)
    print(OUTPUT)
    print("size:", OUTPUT.stat().st_size)


if __name__ == "__main__":
    convert()
